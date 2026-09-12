from collections.abc import AsyncIterator
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient
from pwdlib import PasswordHash
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

from teacher_workspace.config import Settings, get_settings
from teacher_workspace.db import get_session
from teacher_workspace.main import app
from teacher_workspace.models import Base, User
from teacher_workspace.phase7_service import (
    MaterialExtractionError,
    chunk_sections,
    detect_material_type,
    extract_sections,
)
from teacher_workspace.worker import run_once


@pytest.fixture
async def phase7_context(
    tmp_path: Path,
) -> AsyncIterator[tuple[AsyncClient, async_sessionmaker[AsyncSession], Settings]]:
    engine = create_async_engine(
        "sqlite+aiosqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    async with factory() as session, session.begin():
        session.add(
            User(
                username="phase7-teacher",
                password_hash=PasswordHash.recommended().hash("fictional-password"),
                ai_access_enabled=True,
                ai_monthly_job_limit=100,
            )
        )
    settings = Settings(
        database_url="sqlite+aiosqlite://",
        session_secret="phase7-test-secret-that-is-long-enough",
        trusted_origins=["http://test"],
        local_storage_root=tmp_path / "storage",
        ai_provider="mock",
        worker_id="phase7-worker",
    )

    async def override_session() -> AsyncIterator[AsyncSession]:
        async with factory() as session:
            yield session

    app.dependency_overrides[get_session] = override_session
    app.dependency_overrides[get_settings] = lambda: settings
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        yield client, factory, settings
    app.dependency_overrides.clear()
    await engine.dispose()


async def login(client: AsyncClient) -> dict[str, str]:
    response = await client.post(
        "/api/v1/auth/login",
        json={"username": "phase7-teacher", "password": "fictional-password"},
        headers={"Origin": "http://test"},
    )
    assert response.status_code == 200
    csrf = client.cookies.get("teacher_workspace_session_csrf")
    assert csrf
    return {"Origin": "http://test", "X-CSRF-Token": csrf}


async def subject_and_lesson(client: AsyncClient, headers: dict[str, str]) -> tuple[str, str]:
    subject = (
        await client.post("/api/v1/subjects", json={"name": "虚构资料数学"}, headers=headers)
    ).json()
    student = (
        await client.post(
            "/api/v1/students",
            json={"display_name": "资料示例学生", "grade": "七年级"},
            headers=headers,
        )
    ).json()
    link = (
        await client.post(
            "/api/v1/student-subjects",
            json={"student_id": student["id"], "subject_id": subject["id"]},
            headers=headers,
        )
    ).json()
    lesson = (
        await client.post(
            "/api/v1/lessons",
            json={
                "student_subject_id": link["id"],
                "scheduled_start": (datetime.now(UTC) + timedelta(days=1)).isoformat(),
                "planned_minutes": 90,
                "lesson_type": "NEW_LESSON",
                "theme": "虚构整式加减",
            },
            headers=headers,
        )
    ).json()
    return str(link["id"]), str(lesson["id"])


def test_material_extractors_validate_and_chunk() -> None:
    mime, extension = detect_material_type("讲义.txt", "text/plain", "知识点\n例题".encode())
    assert (mime, extension) == ("text/plain", ".txt")
    chunks = chunk_sections(extract_sections(mime, "知识点\n例题".encode()))
    assert chunks[0].text == "知识点\n例题"

    document = Document()
    document.add_heading("虚构讲义", level=1)
    document.add_paragraph("整式加减先合并同类项。")
    output = BytesIO()
    document.save(output)
    docx_mime, _ = detect_material_type(
        "讲义.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        output.getvalue(),
    )
    assert "合并同类项" in extract_sections(docx_mime, output.getvalue())[0].text

    with pytest.raises(MaterialExtractionError):
        detect_material_type("伪装.pdf", "application/pdf", b"not-a-pdf")


@pytest.mark.asyncio
async def test_material_upload_selection_generation_and_archive(
    phase7_context: tuple[AsyncClient, async_sessionmaker[AsyncSession], Settings],
) -> None:
    client, factory, settings = phase7_context
    headers = await login(client)
    link_id, lesson_id = await subject_and_lesson(client, headers)
    uploaded = await client.post(
        "/api/v1/materials",
        data={"student_subject_id": link_id, "purpose": "TEACHING_MATERIAL"},
        files={"file": ("虚构讲义.txt", "整式加减先合并同类项。".encode(), "text/plain")},
        headers=headers,
    )
    assert uploaded.status_code == 201, uploaded.text
    material = uploaded.json()
    assert material["chunk_count"] == 1
    assert material["extracted_chars"] > 0

    usage = await client.get("/api/v1/auth/storage-usage")
    assert usage.status_code == 200
    assert usage.json()["used_bytes"] == len("整式加减先合并同类项。".encode())

    settings.user_upload_quota_bytes = usage.json()["used_bytes"]
    full = await client.post(
        "/api/v1/materials",
        data={"student_subject_id": link_id, "purpose": "TEACHING_MATERIAL"},
        files={"file": ("另一份.txt", "不同的虚构内容".encode(), "text/plain")},
        headers=headers,
    )
    assert full.status_code == 413
    assert full.json()["code"] == "USER_STORAGE_QUOTA_REACHED"

    duplicate = await client.post(
        "/api/v1/materials",
        data={"student_subject_id": link_id, "purpose": "TEACHING_MATERIAL"},
        files={"file": ("重复.txt", "整式加减先合并同类项。".encode(), "text/plain")},
        headers=headers,
    )
    assert duplicate.status_code == 409

    queued = await client.post(
        f"/api/v1/lessons/{lesson_id}/documents/generate",
        json={"material_ids": [material["id"]], "extra_requirements": "引用讲义"},
        headers=headers,
    )
    assert queued.status_code == 202, queued.text
    assert await run_once(factory, settings)
    document = await client.get(f"/api/v1/lessons/{lesson_id}/document")
    assert document.status_code == 200, document.text

    archived = await client.post(
        f"/api/v1/materials/{material['id']}/archive",
        json={"version": material["version"], "reason": "虚构资料测试完成"},
        headers=headers,
    )
    assert archived.status_code == 204
    assert (await client.get("/api/v1/materials")).json() == []
