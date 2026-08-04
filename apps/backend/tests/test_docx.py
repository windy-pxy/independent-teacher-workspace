import io
import zipfile
from datetime import UTC, datetime

import pytest
from docx import Document

from teacher_workspace.docx_generator import (
    LessonDocumentMetadata,
    build_lesson_plan_docx,
    safe_docx_filename,
)
from teacher_workspace.phase2_schemas import LessonPlanContent
from teacher_workspace.prompts import lesson_plan_json_schema
from teacher_workspace.providers.ai import AIRequest, MockAIProvider


def test_lesson_plan_schema_is_strict_for_provider_structured_output() -> None:
    schema = lesson_plan_json_schema()

    assert schema["additionalProperties"] is False
    assert set(schema["required"]) == {
        "schema_version",
        "objectives",
        "schedule",
        "knowledge_explanations",
        "examples",
        "in_class_exercises",
        "common_mistakes",
        "homework",
        "teacher_notes",
    }
    for definition in schema["$defs"].values():
        assert definition["additionalProperties"] is False
        assert set(definition["required"]) == set(definition["properties"])


@pytest.mark.asyncio
async def test_teacher_docx_contains_sections_answers_and_fixed_tables() -> None:
    result = await MockAIProvider().generate(
        AIRequest(
            prompt="虚构教案",
            schema=lesson_plan_json_schema(),
            context={"planned_minutes": 90},
        )
    )
    content = LessonPlanContent.model_validate(result.structured)
    metadata = LessonDocumentMetadata(
        student_alias="示例学生丙",
        grade="八年级",
        subject="物理",
        theme="速度与路程",
        scheduled_start=datetime(2026, 8, 10, 10, 0, tzinfo=UTC),
        planned_minutes=90,
    )

    raw = build_lesson_plan_docx(metadata, content)

    assert raw.startswith(b"PK")
    document = Document(io.BytesIO(raw))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "教学目标" in text
    assert "参考答案与解析" in text
    assert "示例学生丙" in "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert 'w:type="dxa"' in xml
    assert "w:tblGrid" in xml
    assert "w:tcW" in xml
    assert safe_docx_filename(metadata) == "2026-08-10_示例学生丙_物理_速度与路程.docx"
