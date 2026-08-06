from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass

from docx import Document
from pypdf import PdfReader

MAX_EXTRACTED_CHARS = 250_000
MAX_PDF_PAGES = 300
MAX_PDF_PAGE_STREAM_BYTES = 20 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_DOCX_ENTRIES = 10_000
CHUNK_CHARS = 1_500
MAX_CHUNKS = 100


class MaterialExtractionError(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedSection:
    locator: str | None
    text: str


def detect_material_type(filename: str, content_type: str | None, data: bytes) -> tuple[str, str]:
    lowered = filename.lower()
    if lowered.endswith(".pdf") and data.startswith(b"%PDF-"):
        expected = "application/pdf"
        extension = ".pdf"
    elif lowered.endswith(".docx") and data.startswith(b"PK"):
        expected = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        extension = ".docx"
    elif lowered.endswith(".txt"):
        if b"\x00" in data[:4096]:
            raise MaterialExtractionError("TXT 文件包含不支持的二进制内容")
        expected = "text/plain"
        extension = ".txt"
    else:
        raise MaterialExtractionError("仅支持 PDF、DOCX 和 TXT 文件")
    if content_type not in {expected, "application/octet-stream"}:
        raise MaterialExtractionError("文件扩展名、MIME 或内容不一致")
    return expected, extension


def _clean_text(value: str) -> str:
    value = value.replace("\x00", " ")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def _extract_pdf(data: bytes) -> list[ExtractedSection]:
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
    except Exception as exc:
        raise MaterialExtractionError("PDF 文件无法解析") from exc
    if reader.is_encrypted:
        raise MaterialExtractionError("暂不支持加密 PDF")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise MaterialExtractionError(f"PDF 页数不能超过 {MAX_PDF_PAGES} 页")
    sections: list[ExtractedSection] = []
    total = 0
    for index, page in enumerate(reader.pages, start=1):
        try:
            contents = page.get_contents()
            if contents is not None and len(contents.get_data()) > MAX_PDF_PAGE_STREAM_BYTES:
                raise MaterialExtractionError(f"PDF 第 {index} 页内容流过大")
            text = _clean_text(page.extract_text() or "")
        except MaterialExtractionError:
            raise
        except Exception as exc:
            raise MaterialExtractionError(f"PDF 第 {index} 页文本提取失败") from exc
        if text:
            total += len(text)
            if total > MAX_EXTRACTED_CHARS:
                raise MaterialExtractionError("可提取文本超过 250000 字符限制")
            sections.append(ExtractedSection(f"第 {index} 页", text))
    return sections


def _extract_docx(data: bytes) -> list[ExtractedSection]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ENTRIES:
                raise MaterialExtractionError("DOCX 压缩包条目过多")
            if sum(item.file_size for item in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise MaterialExtractionError("DOCX 解压后内容过大")
            names = {item.filename for item in entries}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise MaterialExtractionError("DOCX 文件结构不完整")
        document = Document(io.BytesIO(data))
    except MaterialExtractionError:
        raise
    except Exception as exc:
        raise MaterialExtractionError("DOCX 文件无法解析") from exc
    paragraphs = [_clean_text(item.text) for item in document.paragraphs]
    text = "\n".join(item for item in paragraphs if item)
    for table in document.tables:
        rows = [" | ".join(_clean_text(cell.text) for cell in row.cells) for row in table.rows]
        text += "\n" + "\n".join(item for item in rows if item.strip(" |"))
    text = _clean_text(text)
    if len(text) > MAX_EXTRACTED_CHARS:
        raise MaterialExtractionError("可提取文本超过 250000 字符限制")
    return [ExtractedSection("正文", text)] if text else []


def _extract_txt(data: bytes) -> list[ExtractedSection]:
    try:
        text = _clean_text(data.decode("utf-8-sig"))
    except UnicodeDecodeError as exc:
        raise MaterialExtractionError("TXT 必须使用 UTF-8 编码") from exc
    if len(text) > MAX_EXTRACTED_CHARS:
        raise MaterialExtractionError("可提取文本超过 250000 字符限制")
    return [ExtractedSection("正文", text)] if text else []


def extract_sections(mime_type: str, data: bytes) -> list[ExtractedSection]:
    if mime_type == "application/pdf":
        sections = _extract_pdf(data)
    elif mime_type.endswith("wordprocessingml.document"):
        sections = _extract_docx(data)
    elif mime_type == "text/plain":
        sections = _extract_txt(data)
    else:
        raise MaterialExtractionError("不支持的资料类型")
    if not sections:
        raise MaterialExtractionError("文件没有可提取文本；扫描版 PDF 暂不支持 OCR")
    return sections


def chunk_sections(sections: list[ExtractedSection]) -> list[ExtractedSection]:
    chunks: list[ExtractedSection] = []
    for section in sections:
        paragraphs = [item.strip() for item in section.text.splitlines() if item.strip()]
        current = ""
        for paragraph in paragraphs:
            remaining = paragraph
            while remaining:
                available = CHUNK_CHARS - len(current) - (1 if current else 0)
                if available <= 0:
                    chunks.append(ExtractedSection(section.locator, current))
                    current = ""
                    continue
                piece, remaining = remaining[:available], remaining[available:]
                current = f"{current}\n{piece}" if current else piece
                if len(current) >= CHUNK_CHARS:
                    chunks.append(ExtractedSection(section.locator, current))
                    current = ""
            if len(chunks) >= MAX_CHUNKS:
                raise MaterialExtractionError("资料分段超过允许上限")
        if current:
            chunks.append(ExtractedSection(section.locator, current))
        if len(chunks) > MAX_CHUNKS:
            raise MaterialExtractionError("资料分段超过允许上限")
    return chunks
