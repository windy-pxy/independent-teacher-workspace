from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import datetime
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from teacher_workspace.lesson_plan_contract import LessonPlanContent, QuestionItem


@dataclass(frozen=True)
class LessonDocumentMetadata:
    student_alias: str
    grade: str
    subject: str
    theme: str
    scheduled_start: datetime
    planned_minutes: int


@dataclass(frozen=True)
class TemplateProfile:
    key: str
    accent: str
    soft_fill: str


def profile_for_grade(grade: str) -> TemplateProfile:
    if any(label in grade for label in ("一", "二", "三", "四", "五", "六", "小学")):
        return TemplateProfile("PRIMARY", "2F7D6D", "EAF5F2")
    if any(label in grade for label in ("高一", "高二", "高三", "高中")):
        return TemplateProfile("HIGH", "744A8B", "F2ECF5")
    return TemplateProfile("MIDDLE", "2E5E8C", "E8EEF5")


def _set_cell_margins(
    cell: Any, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120
) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table: Any, widths: list[int], indent: int = 120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)


def _shade(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.first_child_found_in("w:shd")
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_run_font(
    run: Any,
    *,
    latin: str = "Calibri",
    east_asia: str = "宋体",
    size: float = 11,
    bold: bool | None = None,
    color: str | None = None,
) -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _add_markdown_paragraph(doc: DocumentType, text: str, *, style: str | None = None) -> None:
    lines = text.replace("\r\n", "\n").split("\n")
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        is_bullet = line.startswith(("- ", "* "))
        is_number = bool(re.match(r"^\d+[.)]\s+", line))
        if is_bullet:
            line = line[2:].strip()
        elif is_number:
            line = re.sub(r"^\d+[.)]\s+", "", line)
        list_style = "List Bullet" if is_bullet else "List Number" if is_number else None
        paragraph = doc.add_paragraph(style=style or list_style)
        paragraph.paragraph_format.space_after = Pt(6)
        parts = re.split(r"(\*\*[^*]+\*\*)", line)
        for part in parts:
            if not part:
                continue
            bold = part.startswith("**") and part.endswith("**")
            content = part[2:-2] if bold else part.replace("*", "")
            run = paragraph.add_run(content)
            _set_run_font(run, bold=bold)


def _configure_styles(doc: DocumentType, profile: TemplateProfile) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, before, after in (
        ("Heading 1", 16, 18, 10),
        ("Heading 2", 13, 14, 7),
        ("Heading 3", 12, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(profile.accent)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def _add_footer(section: Any, metadata: LessonDocumentMetadata) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(f"{metadata.student_alias} · {metadata.subject} · 教师版  |  第 ")
    _set_run_font(run, size=9, color="666666")
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend((field_begin, instruction, field_end))
    tail = paragraph.add_run(" 页")
    _set_run_font(tail, size=9, color="666666")


def _add_question_group(doc: DocumentType, title: str, questions: list[QuestionItem]) -> None:
    doc.add_heading(title, level=1)
    if not questions:
        _add_markdown_paragraph(doc, "本节未设置该类题目。")
        return
    for index, question in enumerate(questions, 1):
        heading = doc.add_paragraph()
        heading.paragraph_format.keep_with_next = True
        run = heading.add_run(f"{index}. 题目（{question.difficulty}）")
        _set_run_font(run, east_asia="微软雅黑", bold=True)
        _add_markdown_paragraph(doc, question.stem_markdown)


def _add_answer_group(
    doc: DocumentType, title: str, question_groups: list[tuple[str, list[QuestionItem]]]
) -> None:
    doc.add_heading(title, level=1)
    for group_title, questions in question_groups:
        for index, question in enumerate(questions, 1):
            heading = doc.add_paragraph()
            heading.paragraph_format.keep_with_next = True
            run = heading.add_run(f"{group_title} {index}")
            _set_run_font(run, east_asia="微软雅黑", bold=True)
            label = doc.add_paragraph()
            answer_run = label.add_run("答案：")
            _set_run_font(answer_run, bold=True)
            _add_markdown_paragraph(doc, question.answer_markdown)
            analysis = doc.add_paragraph()
            analysis_run = analysis.add_run("解析：")
            _set_run_font(analysis_run, bold=True)
            _add_markdown_paragraph(doc, question.analysis_markdown)


def build_lesson_plan_docx(
    metadata: LessonDocumentMetadata, content: LessonPlanContent
) -> bytes:
    profile = profile_for_grade(metadata.grade)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    _configure_styles(doc, profile)
    _add_footer(section, metadata)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(0)
    run = kicker.add_run(f"{metadata.grade} · {metadata.subject} · 结构化教案")
    _set_run_font(run, east_asia="微软雅黑", size=10, bold=True, color=profile.accent)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run(metadata.theme)
    _set_run_font(title_run, east_asia="微软雅黑", size=24, bold=True, color="0B2545")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle_run = subtitle.add_run("教师版教学设计（审核通过版本）")
    _set_run_font(subtitle_run, east_asia="微软雅黑", size=12, color="555555")

    metadata_table = doc.add_table(rows=2, cols=3)
    metadata_table.style = "Table Grid"
    values = (
        ("学生", metadata.student_alias),
        ("日期", metadata.scheduled_start.strftime("%Y-%m-%d %H:%M")),
        ("时长", f"{metadata.planned_minutes} 分钟"),
    )
    for column, (label, value) in enumerate(values):
        _shade(metadata_table.cell(0, column), profile.soft_fill)
        label_run = metadata_table.cell(0, column).paragraphs[0].add_run(label)
        _set_run_font(label_run, east_asia="微软雅黑", size=9, bold=True, color=profile.accent)
        value_run = metadata_table.cell(1, column).paragraphs[0].add_run(value)
        _set_run_font(value_run, size=10.5)
    _set_table_geometry(metadata_table, [1800, 3960, 3600])

    doc.add_heading("教学目标", level=1)
    for objective in content.objectives:
        _add_markdown_paragraph(doc, f"- {objective}")

    doc.add_heading("时间安排", level=1)
    schedule = doc.add_table(rows=1, cols=3)
    schedule.style = "Table Grid"
    for cell, label in zip(schedule.rows[0].cells, ("时间", "环节", "教学活动"), strict=True):
        _shade(cell, profile.soft_fill)
        run = cell.paragraphs[0].add_run(label)
        _set_run_font(run, east_asia="微软雅黑", size=10, bold=True, color=profile.accent)
    for block in content.schedule:
        cells = schedule.add_row().cells
        cells[0].text = f"{block.minutes} 分钟"
        cells[1].text = block.title
        cells[2].text = re.sub(r"[*_`]", "", block.activities_markdown)
        for cell in cells:
            for run in cell.paragraphs[0].runs:
                _set_run_font(run, size=10)
    _set_table_geometry(schedule, [1440, 2520, 5400])

    doc.add_heading("知识点讲解", level=1)
    for item in content.knowledge_explanations:
        doc.add_heading(item.title, level=2)
        _add_markdown_paragraph(doc, item.body_markdown)

    _add_question_group(doc, "典型例题", content.examples)
    _add_question_group(doc, "当堂练习", content.in_class_exercises)

    doc.add_heading("易错点提醒", level=1)
    for mistake in content.common_mistakes:
        _add_markdown_paragraph(doc, f"- {mistake}")

    _add_question_group(doc, "课后作业", content.homework)

    doc.add_page_break()
    _add_answer_group(
        doc,
        "参考答案与解析",
        [
            ("典型例题", content.examples),
            ("当堂练习", content.in_class_exercises),
            ("课后作业", content.homework),
        ],
    )

    doc.add_heading("教师注意事项", level=1)
    for note in content.teacher_notes:
        _add_markdown_paragraph(doc, f"- {note}")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def safe_docx_filename(metadata: LessonDocumentMetadata) -> str:
    raw = (
        f"{metadata.scheduled_start:%Y-%m-%d}_{metadata.student_alias}_"
        f"{metadata.subject}_{metadata.theme}.docx"
    )
    return re.sub(r"[\\/:*?\"<>|\x00-\x1f]", "_", raw)[:180]
